"""SZSE official XLSX list, bracketed by the same report's dated JSON metadata.

One workbook is one delivery unit, not a fabricated JSON page. Its frozen
bundle retains all three original responses and the source's JSON page count.
"""
from __future__ import annotations

from datetime import date, datetime
import hashlib
from html import unescape
from io import BytesIO
import json
from pathlib import Path
import re
import time
from urllib.parse import urlencode
import warnings
import zipfile
from xml.etree.ElementTree import ParseError

from ..operations.backup import _io
from ..operations.paths import resolve_archived_path
from .exchange_universe import SZSE_ENDPOINT, TZ, szse_page, normalize_szse

ENDPOINT = "https://www.szse.cn/api/report/ShowReport"
PARAMS = {"SHOWTYPE": "xlsx", "CATALOGID": "1110", "TABKEY": "tab1", "random": "0.6935816432433362"}
META_PARAMS = {"SHOWTYPE": "JSON", "CATALOGID": "1110", "TABKEY": "tab1", "PAGENO": 1}
SCHEMA = "szse-export-bundle-v1"
FIELDS = {"板块": "bk", "A股代码": "agdm", "A股简称": "agjc", "A股上市日期": "agssrq", "所属行业": "sshymc"}


def _clean(value):
    return unescape(re.sub(r"<[^>]*>", "", str(value or ""))).strip()


def _normalized(row):
    result = {key: _clean(row.get(key)) for key in FIELDS.values()}
    result["agssrq"] = date.fromisoformat(result["agssrq"]).isoformat()
    return result


def _json(raw):
    def pairs(items):
        value = {}
        for key, item in items:
            if key in value:
                raise ValueError("szse_export_duplicate_json_key")
            value[key] = item
        return value
    def invalid(_):
        raise ValueError("szse_export_nonfinite_json")
    return json.loads(raw, object_pairs_hook=pairs, parse_constant=invalid)


def parse_workbook(raw: bytes) -> list[dict]:
    try:
        return _parse_workbook(raw)
    except (zipfile.BadZipFile, StopIteration, KeyError, IndexError, TypeError, ParseError, OSError) as exc:
        raise ValueError("szse_export_workbook_invalid:" + type(exc).__name__) from exc


def _parse_workbook(raw: bytes) -> list[dict]:
    import openpyxl
    if len(raw) > 4_000_000 or not raw.startswith(b"PK"):
        raise ValueError("szse_export_not_bounded_xlsx")
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 100 or len({item.filename for item in entries}) != len(entries) or sum(item.file_size for item in entries) > 32_000_000:
            raise ValueError("szse_export_expanded_size_limit")
        for item in entries:
            if item.flag_bits & 1 or "externalLinks" in item.filename or "vbaProject" in item.filename:
                raise ValueError("szse_export_active_content_not_allowed")
            if item.filename.endswith(".xml"):
                body = archive.read(item)
                if b"<!DOCTYPE" in body or b"<!ENTITY" in body:
                    raise ValueError("szse_export_xml_entity_not_allowed")
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style", category=UserWarning)
        book = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        if book.sheetnames != ["A股列表"]:
            raise ValueError("szse_export_dataset_sheet_mismatch")
        sheet = book.worksheets[0]
        # The real exporter declares A1:A1. Read actual cells, never trust that
        # declaration as either an empty result or a completeness boundary.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        headers = [cell.value for cell in next(iterator)]
        if len(headers) > 64 or any(headers.count(name) != 1 for name in FIELDS):
            raise ValueError("szse_export_column_contract_changed")
        columns = {field: headers.index(name) for name, field in FIELDS.items()}
        rows, seen = [], set()
        for row_index, cells in enumerate(iterator, 1):
            if row_index > 10000:
                raise ValueError("szse_export_row_position_limit")
            values = [cell.value for cell in cells]
            if not any(value is not None for value in values):
                continue
            if len(rows) >= 10000 or len(values) > 64 or any(cell.data_type == "f" for cell in cells):
                raise ValueError("szse_export_rows_or_formula_invalid")
            row = {key: values[index] if index < len(values) else None for key, index in columns.items()}
            if type(row["agdm"]) is int:
                row["agdm"] = str(row["agdm"]).zfill(6)
            if isinstance(row["agssrq"], (datetime, date)):
                row["agssrq"] = row["agssrq"].date().isoformat() if isinstance(row["agssrq"], datetime) else row["agssrq"].isoformat()
            row = _normalized(row)
            if (not re.fullmatch(r"\d{6}", row["agdm"]) or row["agdm"] in seen
                    or row["bk"] not in {"主板", "创业板"} or not row["agjc"] or not row["sshymc"]):
                raise ValueError("szse_export_missing_duplicate_or_invalid_identity")
            seen.add(row["agdm"])
            rows.append(row)
        return rows
    finally:
        book.close()


def _read(root: Path, locator, refs, expected=None):
    if not isinstance(locator, str) or not locator:
        raise ValueError("szse_export_source_path_missing")
    path = resolve_archived_path(locator, anchor=root)
    if not path.is_absolute():
        path = root / path
    if not path.resolve().is_relative_to(root):
        raise ValueError("szse_export_source_outside_project")
    for parent in (path, *path.parents):
        if parent == root.parent:
            break
        if _io(parent).is_symlink() or (hasattr(parent, "is_junction") and _io(parent).is_junction()):
            raise ValueError("szse_export_source_link_forbidden")
    if _io(path).stat().st_size > 4_000_000:
        raise ValueError("szse_export_source_size_limit")
    raw = _io(path).read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    if expected is not None and expected != digest:
        raise ValueError("szse_export_source_hash_mismatch")
    refs[str(path.resolve())] = {"path": str(path.resolve()), "sha256": digest}
    return raw


def read_bundle(root, bundle: dict, *, target: date, mode: str):
    root = Path(root).resolve()
    if (mode not in {"online", "offline_test"} or bundle.get("schema_version") != SCHEMA
            or bundle.get("provider") != "szse" or bundle.get("dataset") != "a_shares"
            or bundle.get("target_date") != target.isoformat() or bundle.get("provenance_mode") != mode
            or not isinstance(bundle.get("responses"), list) or len(bundle["responses"]) != 3):
        raise ValueError("szse_export_bundle_contract_invalid")
    refs, receipts, payloads, stamps = {}, [], [], []
    for index, reference in enumerate(bundle["responses"]):
        receipt = _json(_read(root, reference["path"], refs, reference["sha256"]))
        expected_endpoint = ENDPOINT if index == 1 else SZSE_ENDPOINT
        actual_params = receipt.get("params")
        if index == 1:
            valid_params = (isinstance(actual_params, dict) and set(actual_params) == set(PARAMS)
                and all(actual_params.get(k) == v for k,v in PARAMS.items() if k != "random")
                and re.fullmatch(r"0\.[0-9]{1,24}", str(actual_params.get("random", ""))) is not None)
        else:
            valid_params = actual_params == META_PARAMS
        if (receipt.get("ok") is not True or receipt.get("http_status") != 200
                or receipt.get("provenance_mode") != mode
                or receipt.get("verification_kind") != ("live_network" if mode == "online" else "offline_test")
                or receipt.get("target_date") != target.isoformat() or not valid_params
                or receipt.get("url") != expected_endpoint + "?" + urlencode(actual_params)):
            raise ValueError("szse_export_http_provenance_or_filter_invalid")
        stamp = datetime.fromisoformat(receipt["fetched_at"])
        if stamp.utcoffset() is None or stamp > datetime.now(TZ) or stamp.astimezone(TZ).date() != target:
            raise ValueError("szse_export_observation_date_invalid")
        if mode == "offline_test" and any("research" in {p.lower() for p in Path(receipt[key]).parts} for key in ["raw_path", "response_path"]):
            raise ValueError("szse_export_test_evidence_in_research")
        raw = _read(root, receipt["raw_path"], refs, receipt["raw_sha256"])
        if index == 1:
            if receipt.get("body_format") != "xlsx":
                raise ValueError("szse_export_format_invalid")
            payload = parse_workbook(raw)
        else:
            payload = _json(raw.decode("utf-8-sig"))
            if payload != receipt.get("body"):
                raise ValueError("szse_export_metadata_raw_conflict")
        receipts.append(receipt); payloads.append(payload); stamps.append(stamp)
    if not stamps[0] <= stamps[1] <= stamps[2] or (stamps[2] - stamps[0]).total_seconds() > 180:
        raise ValueError("szse_export_bracketing_time_invalid")
    first, before = szse_page(payloads[0], requested_page=1)
    last, after = szse_page(payloads[2], requested_page=1)
    meta = next(tab["metadata"] for tab in payloads[0] if tab.get("metadata", {}).get("tabkey") == "tab1")
    final_meta = next(tab["metadata"] for tab in payloads[2] if tab.get("metadata", {}).get("tabkey") == "tab1")
    if (before != after or first != last or meta != final_meta or before["source_as_of_date"] != target.isoformat()
            or len(payloads[1]) != before["expected_records"] or not payloads[1]):
        raise ValueError("szse_export_count_date_or_boundary_conflict")
    by_code = {row["agdm"]: row for row in payloads[1]}
    if any(by_code.get(str(row.get("agdm"))) != _normalized(row) for row in first):
        raise ValueError("szse_export_first_page_mismatch")
    if any(row["agssrq"] > target.isoformat() for row in payloads[1]):
        raise ValueError("szse_export_future_listing_date")
    boundary = {"page_number": 1, "page_size": len(payloads[1]), "expected_pages": 1,
                "expected_records": before["expected_records"], "terminal": True,
                "source_as_of_date": target.isoformat(), "transport": "official_xlsx_export",
                "source_expected_pages": before["expected_pages"]}
    return {"records": payloads[1], "boundary": boundary, "metadata": meta,
            "observed_at": receipts[2]["fetched_at"], "file_refs": list(refs.values())}


def collect_export(fetch, directory: Path, *, target: date, mode: str, requests: list):
    refs = []
    # random is the public export's cache-busting parameter, not an access token.
    # Keep the observed numeric shape but do not copy a constant cache key from
    # a library example into every scheduled collection.
    export_params = {**PARAMS, "random": "0." + str(time.time_ns())}
    for label, endpoint, params in [("before", SZSE_ENDPOINT, META_PARAMS), ("xlsx", ENDPOINT, export_params),
                                     ("after", SZSE_ENDPOINT, META_PARAMS)]:
        response = fetch(endpoint, dict(params), label="szse-export-" + label, target=target)
        requests.append({k:v for k,v in response.items() if k != "body"})
        if response.get("ok") is not True:
            raise ValueError(response.get("status", "szse_export_request_failed"))
        if label != "xlsx":
            _, bounds = szse_page(response["body"], requested_page=1)
            if bounds["source_as_of_date"] != target.isoformat():
                raise ValueError("source_not_updated_for_target_date")
        path = Path(response["response_path"])
        refs.append({"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    bundle = {"schema_version": SCHEMA, "provider": "szse", "dataset": "a_shares",
              "target_date": target.isoformat(), "provenance_mode": mode, "responses": refs}
    validated = read_bundle(directory, bundle, target=target, mode=mode)
    path = directory / "requests/szse-export-bundle.json"
    with path.open("x", encoding="utf-8") as stream:
        json.dump(bundle, stream, ensure_ascii=False, indent=2)
    hashed = hashlib.sha256(path.read_bytes()).hexdigest()
    part = {"provider": "szse", "dataset": "a_shares", "page_number": 1,
            "records": [normalize_szse(row, evidence_id=hashed) for row in validated["records"]],
            "terminal": True, "source_boundary": validated["boundary"],
            "raw_response_path": str(path), "raw_sha256": hashed}
    updates = {"expected_pages": 1, "expected_records": len(validated["records"]), "complete": True,
               "source_expected_pages": validated["boundary"]["source_expected_pages"],
               "source_transport": "official_xlsx_export", "as_of_date": target.isoformat(),
               "source_as_of_date": target.isoformat(), "source_business_date": target.isoformat(),
               "observed_at": validated["observed_at"], "temporal_basis": "dated_metadata_bracketed_current_export"}
    return part, updates
